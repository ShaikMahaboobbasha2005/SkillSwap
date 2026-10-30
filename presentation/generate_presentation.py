import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # SkillSwap Color Palette (from Design.md)
    PINE_DARK   = RGBColor(27, 67, 50)      # #1B4332 Primary Brand
    PINE_LIGHT  = RGBColor(228, 238, 232)   # #E4EEE8 Soft Accent
    EMERALD     = RGBColor(63, 168, 115)    # #3FA873 Accent Active
    BG_CREAM    = RGBColor(247, 246, 242)   # #F7F6F2 Main Light BG
    SURFACE_WHITE = RGBColor(255, 255, 255) # #FFFFFF Cards
    INK_MAIN    = RGBColor(22, 22, 15)      # #16160F Primary Text
    INK_MUTED   = RGBColor(107, 104, 88)    # #6B6858 Secondary Text
    BORDER_COL  = RGBColor(230, 227, 218)   # #E6E3DA Dividers/Borders
    CHARCOAL_BG = RGBColor(15, 18, 16)      # #0F1210 Dark Mode / Header
    GOLD        = RGBColor(184, 134, 11)    # #B8860B Rating Stars
    CARD_BG_ALT = RGBColor(240, 244, 241)   # Slightly tinted card bg

    # Assets
    base_dir = r"C:\SkillSwap"
    logo_path = os.path.join(base_dir, "client", "src", "assets", "logo.png")
    hero_path = os.path.join(base_dir, "client", "src", "assets", "hero.png")
    icon_matches = os.path.join(base_dir, "client", "src", "assets", "icons", "matches.png")
    icon_offered = os.path.join(base_dir, "client", "src", "assets", "icons", "offered-skill.png")
    icon_wanted = os.path.join(base_dir, "client", "src", "assets", "icons", "wanted-skill.png")

    def add_slide_header(slide, title, category="SKILLSWAP — BTECH MINI PROJECT"):
        # Header background banner
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = PINE_DARK
        header_box.line.color.rgb = PINE_DARK

        # Category eyebrow
        tf_cat = header_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = Inches(0.8)
        tf_cat.margin_top = Inches(0.18)
        tf_cat.margin_right = Inches(0.8)
        tf_cat.margin_bottom = Inches(0)
        
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = EMERALD
        p_cat.font.name = "Calibri"

        # Main Slide Title
        p_title = tf_cat.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = SURFACE_WHITE
        p_title.font.name = "Calibri"

        # Bottom accent line on header
        accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.04))
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = EMERALD
        accent_line.line.color.rgb = EMERALD

        # Slide canvas background
        bg_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.19), Inches(13.333), Inches(6.31))
        bg_rect.fill.solid()
        bg_rect.fill.fore_color.rgb = BG_CREAM
        bg_rect.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=SURFACE_WHITE, border_color=BORDER_COL):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    # Dark charcoal luxury backdrop
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = CHARCOAL_BG
    bg1.line.fill.background()

    # Emerald brand accent bar on top
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = EMERALD
    top_bar.line.fill.background()

    # Left side: Brand, Title, Subtitle
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(1.2), Inches(1.2), width=Inches(3.2))

    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.7), Inches(7.2), Inches(3.8))
    tf = t_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "B.TECH MINI PROJECT PRESENTATION"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = EMERALD
    p0.font.name = "Calibri"

    p1 = tf.add_paragraph()
    p1.text = "SkillSwap"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = SURFACE_WHITE
    p1.font.name = "Georgia"
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "A Peer-to-Peer Skill Exchange Platform"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(200, 220, 210)
    p2.font.name = "Calibri"
    p2.space_before = Pt(4)

    p3 = tf.add_paragraph()
    p3.text = "Share What You Know. Master What You Do.\nA modern full-stack barter network replacing expensive courses with reciprocal knowledge trades."
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(160, 175, 168)
    p3.font.name = "Calibri"
    p3.space_before = Pt(14)

    # Right Card: Project & Team Details
    team_card = add_card(slide1, Inches(8.8), Inches(1.2), Inches(3.7), Inches(5.3), bg_color=RGBColor(24, 28, 26), border_color=PINE_DARK)
    t_team = team_card.text_frame
    t_team.word_wrap = True
    t_team.margin_left = Inches(0.35)
    t_team.margin_top = Inches(0.4)
    t_team.margin_right = Inches(0.35)

    pt1 = t_team.paragraphs[0]
    pt1.text = "PROJECT TEAM"
    pt1.font.size = Pt(12)
    pt1.font.bold = True
    pt1.font.color.rgb = EMERALD

    members = [
        ("Shaikh Mahaboob Basha", "Lead Full-Stack Developer"),
        ("Shaik Mohammad Juber", "Team Member"),
        ("Shaik DadaPeer", "Team Member"),
        ("Shaik Thousif", "Team Member")
    ]
    for name, role in members:
        pm_name = t_team.add_paragraph()
        pm_name.text = f"• {name}"
        pm_name.font.size = Pt(13)
        pm_name.font.bold = True
        pm_name.font.color.rgb = SURFACE_WHITE
        pm_name.space_before = Pt(10)

        pm_role = t_team.add_paragraph()
        pm_role.text = f"   {role}"
        pm_role.font.size = Pt(11)
        pm_role.font.color.rgb = RGBColor(150, 170, 160)

    pt_tech = t_team.add_paragraph()
    pt_tech.text = "CORE TECHNOLOGIES"
    pt_tech.font.size = Pt(11)
    pt_tech.font.bold = True
    pt_tech.font.color.rgb = EMERALD
    pt_tech.space_before = Pt(20)

    pt_tech_val = t_team.add_paragraph()
    pt_tech_val.text = "React 19 • Node.js • Express v5 • MongoDB Atlas • Socket.IO • Google Gemini AI • Jitsi Meet"
    pt_tech_val.font.size = Pt(10)
    pt_tech_val.font.color.rgb = RGBColor(180, 190, 185)
    pt_tech_val.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: ABSTRACT
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide2, "Abstract — Executive Overview", "EXECUTIVE SUMMARY")

    # 3 Main Highlights Cards
    cards_data2 = [
        ("The Core Concept", "Bilateral Knowledge Barter", 
         "SkillSwap is a production-grade web application enabling users to trade skills directly with one another without fiat currency. Members list skills they can teach ('Offered') and skills they wish to learn ('Wanted'), unlocking mutual skill reciprocity."),
        ("Key Architecture", "Modern Full-Stack Stack", 
         "Engineered with a React 19 SPA, Express.js v5 REST API, and MongoDB Atlas. Features real-time Socket.IO room messaging, embedded WebRTC video meetings via Jitsi, Cloudinary media storage, and Google Gemini Flash AI matching."),
        ("Accountability & Trust", "Two-Party Verification", 
         "Eliminates trust deficits via strict two-party completion confirmations, verified star ratings & reviews linked to exact skill pairs, and an Instagram-style portfolio gallery with 4 curated reactions.")
    ]
    for i, (tag, title, desc) in enumerate(cards_data2):
        x = Inches(0.8 + i * 4.0)
        c = add_card(slide2, x, Inches(1.6), Inches(3.7), Inches(3.8))
        tf_c = c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.3)
        tf_c.margin_top = Inches(0.35)
        tf_c.margin_right = Inches(0.3)

        p = tf_c.paragraphs[0]
        p.text = tag.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = PINE_DARK

        p2 = tf_c.add_paragraph()
        p2.text = title
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = INK_MAIN
        p2.space_before = Pt(6)

        p3 = tf_c.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(12)
        p3.font.color.rgb = INK_MUTED
        p3.space_before = Pt(10)

    # Bottom Summary Banner
    banner2 = add_card(slide2, Inches(0.8), Inches(5.7), Inches(11.7), Inches(1.2), bg_color=PINE_LIGHT, border_color=EMERALD)
    tb2 = banner2.text_frame
    tb2.word_wrap = True
    tb2.margin_left = Inches(0.4)
    tb2.margin_top = Inches(0.2)
    p_ban = tb2.paragraphs[0]
    p_ban.text = "🎯 PROJECT MISSION & VALUE PROPOSITION"
    p_ban.font.size = Pt(11)
    p_ban.font.bold = True
    p_ban.font.color.rgb = PINE_DARK
    p_ban2 = tb2.add_paragraph()
    p_ban2.text = "SkillSwap democratizes lifelong learning by transforming hidden individual competencies into shared community capital. By anchoring recommendations to deterministic skill overlaps and augmenting with AI semantic reasoning, it delivers authentic, verifiable peer learning."
    p_ban2.font.size = Pt(11)
    p_ban2.font.color.rgb = INK_MAIN
    p_ban2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide3, "Problem Statement — Challenges in Skill Learning", "PROBLEM DEFINITION")

    problems = [
        ("💸 Financial Barriers in Education",
         "Online courses (Coursera, Udemy) and private tutors charge steep recurring subscription or hourly fees. High costs prevent economically diverse learners from mastering high-value crafts."),
        ("📦 Latent & Underutilized Expertise",
         "Millions of individuals possess practical skills (e.g., frontend coding, design, guitar, languages) but have no direct outlet to exchange their knowledge for topics they urgently need to learn."),
        ("🤝 Trust Deficits & Lack of Accountability",
         "Conventional social platforms offer no structured verification. Without two-party verification, verified ratings, and work portfolios, users cannot evaluate the authenticity or credibility of peer instructors."),
        ("🔀 Tool Fragmentation & Coordination Friction",
         "Traditional informal exchanges force users to juggle WhatsApp for messaging, Zoom for calling, Google Calendar for scheduling, and third-party links for payments, leading to severe communication breakdown.")
    ]

    for i, (title, desc) in enumerate(problems):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(1.6 + row * 2.7)
        c = add_card(slide3, x, y, Inches(5.7), Inches(2.4))
        tf_c = c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.35)
        tf_c.margin_top = Inches(0.3)
        tf_c.margin_right = Inches(0.35)

        p = tf_c.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = PINE_DARK

        p2 = tf_c.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = INK_MUTED
        p2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 4: OBJECTIVES
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide4, "Project Objectives — Goals of SkillSwap", "PROJECT SPECIFICATIONS")

    objectives = [
        ("1. Peer-to-Peer Barter Engine", "Build a symmetric exchange framework where users manage Offered and Wanted skills, enabling direct barter without fiat currency."),
        ("2. Hybrid AI Skill Matching", "Develop a two-tier recommendation engine combining deterministic progressive diminishing-returns scoring with Google Gemini Flash semantic re-ranking."),
        ("3. Unified Communication Hub", "Integrate real-time Socket.IO chat (with sent/delivered/read receipts) and in-browser WebRTC private video calling via embedded Jitsi Meet."),
        ("4. Trust & Reputation Architecture", "Enforce two-party completion confirmations before ratings unlock; dynamically aggregate server-managed reputation scores exclusively from verified swaps."),
        ("5. Portfolio Media Showcase", "Implement an Instagram-style media gallery supporting signed Cloudinary image and video uploads, emoji reactions, and 5-second atomic undo protection."),
        ("6. Production-Grade UI/UX", "Deliver a responsive Single Page App featuring scroll-driven 3D WebGL scenes (Three.js), complete System/Light/Dark mode with zero FOUC, and accessible mobile bottom navigation.")
    ]

    for i, (title, desc) in enumerate(objectives):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(1.6 + row * 1.8)
        c = add_card(slide4, x, y, Inches(5.7), Inches(1.6))
        tf_c = c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.3)
        tf_c.margin_top = Inches(0.2)
        tf_c.margin_right = Inches(0.3)

        p = tf_c.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = PINE_DARK

        p2 = tf_c.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = INK_MUTED
        p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 5: PROPOSED SOLUTION
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide5, "Proposed Solution — How SkillSwap Works", "SOLUTION ARCHITECTURE")

    # Workflow columns: 4 pillars
    pillars = [
        ("01. Profile & Taxonomy", "Skills Offered & Wanted", 
         "Users define up to 5 teaching skills ('Offer') and learning goals ('Learn') categorized across 20 disciplines with defined proficiency levels."),
        ("02. Intelligent Discovery", "Deterministic + Gemini AI", 
         "A progressive compatibility engine computes exact overlaps and category affinities; Gemini Flash evaluates domain relationships under a 70/30 hybrid anchor."),
        ("03. Real-Time Workspace", "Live Chat & Video Calling", 
         "Accepted swaps open dedicated Socket.IO chat rooms with delivery receipts and integrated one-click Jitsi Meet video conferencing without leaving the app."),
        ("04. Verified Completion", "Two-Party Rating & Reviews", 
         "Both parties must mutually confirm completion to finalize the swap, archiving conversations and unlocking verified 1-5 star reviews that drive reputation.")
    ]

    for i, (step, title, desc) in enumerate(pillars):
        x = Inches(0.8 + i * 3.0)
        c = add_card(slide5, x, Inches(1.6), Inches(2.75), Inches(3.8))
        tf_c = c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_top = Inches(0.3)
        tf_c.margin_right = Inches(0.25)

        p = tf_c.paragraphs[0]
        p.text = step
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = EMERALD

        p2 = tf_c.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = INK_MAIN
        p2.space_before = Pt(6)

        p3 = tf_c.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = INK_MUTED
        p3.space_before = Pt(10)

    # Solution Advantage comparison bar
    comp_box = add_card(slide5, Inches(0.8), Inches(5.65), Inches(11.7), Inches(1.3), bg_color=PINE_LIGHT, border_color=PINE_DARK)
    tfc = comp_box.text_frame
    tfc.word_wrap = True
    tfc.margin_left = Inches(0.3)
    tfc.margin_top = Inches(0.18)
    pc = tfc.paragraphs[0]
    pc.text = "✨ KEY DISTINCTION: EXISTING MARKETPLACE VS. SKILLSWAP"
    pc.font.size = Pt(11)
    pc.font.bold = True
    pc.font.color.rgb = PINE_DARK

    pc2 = tfc.add_paragraph()
    pc2.text = "• Traditional Platforms: One-way transaction ($20–$100/hr) • Third-party fees • Fragmented tools • Unverified reviews\n• SkillSwap: Bilateral barter (0 Currency) • End-to-end integrated video & chat • Mutual completion verification • Real portfolio proof"
    pc2.font.size = Pt(11)
    pc2.font.color.rgb = INK_MAIN
    pc2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 6: TECHNOLOGY STACK
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide6, "Technology Stack — Comprehensive Implementation", "TECH SPECIFICATIONS")

    tech_categories = [
        ("Frontend Client", [
            ("React 19", "Modern UI with React Compiler for automated memoization"),
            ("Vite 8", "High-performance build tooling & instantaneous HMR"),
            ("Tailwind CSS v4", "Next-gen utility styling, CSS-first design tokens"),
            ("Three.js (WebGL)", "Scroll-driven 5-stage 3D landing page scene"),
            ("Socket.IO Client", "WebSocket client for live chat & real-time badge sync")
        ]),
        ("Backend Services", [
            ("Node.js (v20+)", "Asynchronous event-driven server runtime"),
            ("Express.js v5", "REST API routing with native async error handling"),
            ("Socket.IO v4", "Room-based WebSocket server on same HTTP instance"),
            ("Zod & Joi", "Strict schema validation before database access"),
            ("Bcrypt.js & JWT", "Stateless auth, 10-round salted password hashing")
        ]),
        ("Database & Cloud", [
            ("MongoDB Atlas", "Multi-collection hosted NoSQL document database"),
            ("Mongoose v9", "Schema modeling, compound indexing & aggregations"),
            ("Cloudinary SDK", "Signed image & video uploads with auto lifecycle deletion"),
            ("Google Gemini Flash", "Semantic AI matching layer via @google/genai"),
            ("Jitsi Meet API", "Zero-account WebRTC private video conferencing")
        ])
    ]

    for i, (cat_title, items) in enumerate(tech_categories):
        x = Inches(0.8 + i * 4.0)
        c = add_card(slide6, x, Inches(1.6), Inches(3.7), Inches(5.3))
        tf_c = c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.3)
        tf_c.margin_top = Inches(0.3)
        tf_c.margin_right = Inches(0.3)

        p = tf_c.paragraphs[0]
        p.text = cat_title.upper()
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PINE_DARK

        for tech, desc in items:
            pt = tf_c.add_paragraph()
            pt.text = f"• {tech}"
            pt.font.size = Pt(12)
            pt.font.bold = True
            pt.font.color.rgb = INK_MAIN
            pt.space_before = Pt(8)

            pd = tf_c.add_paragraph()
            pd.text = f"   {desc}"
            pd.font.size = Pt(10)
            pd.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 7: SYSTEM ARCHITECTURE
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide7, "System Architecture — Component Topology", "ARCHITECTURE")

    # Left box: Client Layer
    c_client = add_card(slide7, Inches(0.8), Inches(1.6), Inches(3.4), Inches(5.3), bg_color=CARD_BG_ALT)
    tfc = c_client.text_frame
    tfc.word_wrap = True
    tfc.margin_left = Inches(0.3)
    tfc.margin_top = Inches(0.3)
    p = tfc.paragraphs[0]
    p.text = "CLIENT LAYER (SPA)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    client_items = [
        ("React 19 Pages", "13 Route views: Dashboard, Discover, Matches, Swaps, Chats, Profile, etc."),
        ("Context State", "AuthContext, SocketContext, ThemeContext, NotificationContext"),
        ("Service Wrappers", "Axios clients with JWT bearer interceptors"),
        ("3D Engine", "Three.js WebGL canvas for unauthenticated landing page")
    ]
    for h, d in client_items:
        ph = tfc.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = INK_MAIN
        ph.space_before = Pt(8)
        pd = tfc.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = INK_MUTED

    # Middle box: Monolithic Backend Server
    c_server = add_card(slide7, Inches(4.6), Inches(1.6), Inches(4.2), Inches(5.3), bg_color=SURFACE_WHITE, border_color=PINE_DARK)
    tfs = c_server.text_frame
    tfs.word_wrap = True
    tfs.margin_left = Inches(0.3)
    tfs.margin_top = Inches(0.3)
    p = tfs.paragraphs[0]
    p.text = "EXPRESS.JS BACKEND MONOLITH"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    server_items = [
        ("REST API Routes (12)", "/api/auth, /profile, /skills, /discover, /swaps, /chat, /meetings, /ratings, etc."),
        ("Thin Controllers, Fat Services", "Business logic isolated in dedicated testable service modules"),
        ("Socket.IO Engine", "Rooms: 'swap:<swapId>' (chat) & 'user:<userId>' (real-time notifications)"),
        ("Background Jobs", "Atomic 30s scheduler for 15-min and start-time meeting reminders"),
        ("Validation & Auth Guards", "Zod schemas + JWT verification middleware on all protected endpoints")
    ]
    for h, d in server_items:
        ph = tfs.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = INK_MAIN
        ph.space_before = Pt(8)
        pd = tfs.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = INK_MUTED

    # Right box: External Services
    c_ext = add_card(slide7, Inches(9.2), Inches(1.6), Inches(3.3), Inches(5.3), bg_color=CARD_BG_ALT)
    tfe = c_ext.text_frame
    tfe.word_wrap = True
    tfe.margin_left = Inches(0.3)
    tfe.margin_top = Inches(0.3)
    p = tfe.paragraphs[0]
    p.text = "EXTERNAL SERVICES"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    ext_items = [
        ("MongoDB Atlas", "9 Collections with compound unique indexes & aggregation pipelines"),
        ("Cloudinary Cloud Media", "Storage for avatars, banners, and portfolio media with lifecycle cleanup"),
        ("Google Gemini Flash", "Semantic AI matching for taxonomy relations via @google/genai"),
        ("Jitsi Meet WebRTC", "Embedded private video conferencing rooms")
    ]
    for h, d in ext_items:
        ph = tfe.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = INK_MAIN
        ph.space_before = Pt(8)
        pd = tfe.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 8: APPLICATION FLOWCHART
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide8, "Application Workflow — End-to-End User Journey", "PROCESS FLOWCHART")

    workflow_steps = [
        ("1. Onboard", "Sign up with email/password; setup profile with location, bio, and avatar."),
        ("2. Define Skills", "List Offered skills (what I teach) and Wanted skills (what I want to learn)."),
        ("3. Discover / Match", "Browse Discover grid or view AI-ranked Matches with compatibility %."),
        ("4. Swap Proposal", "Select candidate, choose reciprocal skills, and submit swap proposal."),
        ("5. Accept Swap", "Recipient reviews proposal and clicks 'Accept' to transition status."),
        ("6. Chat & Meet", "Real-time chat activates with WhatsApp-style receipts & Jitsi video calls."),
        ("7. Two-Party Done", "Participant clicks 'Mark Completed'; partner confirms completion."),
        ("8. Rate & Review", "Both leave 1-5 star reviews; server recalculates average reputation.")
    ]

    for i, (title, desc) in enumerate(workflow_steps):
        row = i // 4
        col = i % 4
        x = Inches(0.8 + col * 2.95)
        y = Inches(1.6 + row * 2.7)
        c = add_card(slide8, x, y, Inches(2.8), Inches(2.4), bg_color=SURFACE_WHITE)
        tfc = c.text_frame
        tfc.word_wrap = True
        tfc.margin_left = Inches(0.25)
        tfc.margin_top = Inches(0.25)
        tfc.margin_right = Inches(0.25)

        p = tfc.paragraphs[0]
        p.text = f"STEP {i+1}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = EMERALD

        p2 = tfc.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = PINE_DARK
        p2.space_before = Pt(4)

        p3 = tfc.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = INK_MUTED
        p3.space_before = Pt(8)

    # =========================================================================
    # SLIDE 9: KEY FEATURES
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide9, "Key Implemented Features", "CORE CAPABILITIES")

    features = [
        ("🎨 3D WebGL Landing Page", "Cinematic Three.js scroll journey with abstract sculptural avatars, frosted glass skill capsules, and mobile GPU optimization."),
        ("🌓 Zero-FOUC Dark Mode", "System/Light/Dark switching with synchronous inline <head> evaluation preventing any unstyled content flash."),
        ("🔄 Multi-State Swap Machine", "Full state machine (pending -> accepted -> completed, or rejected/cancelled/left) with atomic two-party completion protection."),
        ("💬 Live Socket.IO Chat", "Isolated swap rooms, persistent receipts (sent ✓ -> delivered ✓✓ -> read ✓✓), soft deletion, and date boundaries."),
        ("📹 Integrated Jitsi Meetings", "Instant or scheduled video calls directly inside chat with atomic 30s background reminder worker (15m advance alerts)."),
        ("🖼️ Portfolio Showcase & Undo", "Instagram-style image/video gallery with 4 reactions (👍🔥👏💡), 5-second atomic undo toast, and moderation report flags."),
        ("🔔 Notification Center", "Interactive Navbar bell with live pulse badge, full /notifications page, and deep-link auto-scroll element highlighting."),
        ("📱 Mobile-First Architecture", "5-destination bottom navigation bar, safe-area inset clearance, and proprietary custom SVG icon family.")
    ]

    for i, (title, desc) in enumerate(features):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(1.55 + row * 1.35)
        c = add_card(slide9, x, y, Inches(5.7), Inches(1.22))
        tfc = c.text_frame
        tfc.word_wrap = True
        tfc.margin_left = Inches(0.25)
        tfc.margin_top = Inches(0.18)
        tfc.margin_right = Inches(0.25)

        p = tfc.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PINE_DARK

        p2 = tfc.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = INK_MUTED
        p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 10: AI-BASED SKILL MATCHING
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide10, "Smart Recommendations & Google Gemini AI", "AI ALGORITHMS")

    # Left: Deterministic Matching
    c_det = add_card(slide10, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.3))
    tfd = c_det.text_frame
    tfd.word_wrap = True
    tfd.margin_left = Inches(0.35)
    tfd.margin_top = Inches(0.3)
    tfd.margin_right = Inches(0.35)

    p = tfd.paragraphs[0]
    p.text = "TIER 1: DETERMINISTIC MATCH ENGINE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    det_points = [
        ("Canonical Skill Normalization", "Centralized alias dictionary standardizes spelling (e.g. 'JS' <-> 'JavaScript', 'ReactJS' <-> 'React') while enforcing false-equivalence guards ('Java' != 'JavaScript')."),
        ("Progressive Diminishing Returns Curve", "Directional score formula (0->0, 1->26, 2->36, 3->42, 4->45, 5+->46) prevents premature 100% score saturation. Multi-skill matches differentiate clearly (1<->1 = 72%, 2<->2 = 92%)."),
        ("Mutual Exchange Bonus (+20)", "Two-way complementary exchanges receive a direct +20 bonus and high-visibility badge."),
        ("Related Category Affinities (+8)", "Non-exact skills sharing a category (e.g. Web Dev) earn +8 each (max 16), strictly clamped to 0–100."),
        ("Preserved Display Names", "Human-readable match explanations always cite original user-entered skill names.")
    ]
    for h, d in det_points:
        ph = tfd.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = INK_MAIN
        ph.space_before = Pt(6)
        pd = tfd.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = INK_MUTED

    # Right: Gemini AI Integration
    c_ai = add_card(slide10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3))
    tfa = c_ai.text_frame
    tfa.word_wrap = True
    tfa.margin_left = Inches(0.35)
    tfa.margin_top = Inches(0.3)
    tfa.margin_right = Inches(0.35)

    p = tfa.paragraphs[0]
    p.text = "TIER 2: GOOGLE GEMINI FLASH SEMANTIC LAYER"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    ai_points = [
        ("User-Triggered Opt-In (?ai=true)", "Default recommendations run 100% deterministically with 0 external API calls for instant response. AI is triggered only when explicitly requested."),
        ("Structured Taxonomy Evaluation", "Gemini Flash evaluates deeper cross-discipline affinities (e.g., 'Web Development' <-> 'React' + 'Node.js') for a bounded candidate pool (<= 20 users)."),
        ("70/30 Hybrid Scoring Guardrail", "Formula: hybridScore = round(traditionalScore * 0.70 + semanticScore * 0.30). The 70% deterministic weight ensures true matches cannot be displaced by AI hallucination."),
        ("Privacy-Safe Bounded Prompts", "Passes only sanitized skill names and levels. Zero personal data (no passwords, emails, or chat logs). Candidate IDs verified to prevent fabrication."),
        ("Resilient Fallback", "Gracefully falls back to 100% deterministic ranking if API key is absent, rate-limited (429), or times out, ensuring zero UI crash.")
    ]
    for h, d in ai_points:
        ph = tfa.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = INK_MAIN
        ph.space_before = Pt(6)
        pd = tfa.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 11: DATABASE DESIGN
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide11, "Database Design — MongoDB / Mongoose Schema", "DATABASE ARCHITECTURE")

    collections_data = [
        ("User", "name, email (unique), passwordHash, profilePicture, profileBanner, location, bio, socialLinks, avgRating (server-managed), completedSwaps (server-managed)"),
        ("Skill", "owner (ref: User), name, normalizedName, category (enum 20), level (enum 4), type ('Offer'|'Learn'), status ('Active'|'Inactive'). Unique: {owner, type, normalizedName}"),
        ("SwapRequest", "fromUser, toUser, offeredSkill, wantedSkill, immutable skill snapshots, status ('pending'|'accepted'|'completed'|'rejected'|'cancelled'|'left'), completionRequestedBy"),
        ("Message", "swapRequest (ref), sender (ref), type ('text'|'meeting'), content, meetingSession (ref), isDeleted, status ('sent'|'delivered'|'read'), replyTo"),
        ("MeetingSession", "swap (ref), participants, roomName, scheduledAt, duration (15|30|45|60), status ('scheduled'|'active'|'completed'|'cancelled'), reminded15m, remindedStart"),
        ("Rating", "swapRequest (ref), reviewer (ref), ratedUser (ref), stars (1-5), review. Compound unique index: {swapRequest, reviewer} enforcing max 1 rating per swap"),
        ("Notification", "user (ref), sender (ref), swap (ref), type (12 enums), title, message, read (bool). Compound index: {user, read, createdAt: -1}"),
        ("Portfolio", "user (ref), media {url, publicId, type, duration}, caption, skill (ref), moderationStatus, reactions [{user, type}], reportCount. Max 30 items per user"),
        ("PortfolioReport", "portfolioItem (ref), reporter (ref), reason (7 enums), description, status. Compound unique index: {portfolioItem, reporter}")
    ]

    for i, (name, fields) in enumerate(collections_data):
        row = i // 3
        col = i % 3
        x = Inches(0.8 + col * 4.0)
        y = Inches(1.55 + row * 1.8)
        c = add_card(slide11, x, y, Inches(3.7), Inches(1.68))
        tfc = c.text_frame
        tfc.word_wrap = True
        tfc.margin_left = Inches(0.25)
        tfc.margin_top = Inches(0.2)
        tfc.margin_right = Inches(0.25)

        p = tfc.paragraphs[0]
        p.text = f"📁 {name}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PINE_DARK

        p2 = tfc.add_paragraph()
        p2.text = fields
        p2.font.size = Pt(10)
        p2.font.color.rgb = INK_MUTED
        p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 12: SCREENSHOTS & UI DESIGN
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide12, "User Interface & Verified Design System", "UI & DESIGN SYSTEM")

    # Left: Brand & Iconography Showcase
    c_brand = add_card(slide12, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.3))
    tfb = c_brand.text_frame
    tfb.word_wrap = True
    tfb.margin_left = Inches(0.3)
    tfb.margin_top = Inches(0.3)
    tfb.margin_right = Inches(0.3)

    p = tfb.paragraphs[0]
    p.text = "PROPRIETARY BRAND & ICONS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    if os.path.exists(hero_path):
        slide12.shapes.add_picture(hero_path, Inches(1.1), Inches(2.2), width=Inches(3.8))

    p2 = tfb.add_paragraph()
    p2.text = "Custom SVG Icon Family:"
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = INK_MAIN
    p2.space_before = Pt(140)

    icon_notes = [
        ("📄 OfferedSkillIcon", "Hand presenting document with pencil"),
        ("📑 WantedSkillIcon", "Hand holding sheets with bookmark"),
        ("🪐 MatchesIcon", "Orbital reciprocal connection nodes")
    ]
    for iname, idesc in icon_notes:
        pi = tfb.add_paragraph()
        pi.text = f"• {iname}: {idesc}"
        pi.font.size = Pt(10)
        pi.font.color.rgb = INK_MUTED
        pi.space_before = Pt(4)

    # Right: Verified UI Views Architecture
    c_views = add_card(slide12, Inches(5.6), Inches(1.6), Inches(6.9), Inches(5.3))
    tfv = c_views.text_frame
    tfv.word_wrap = True
    tfv.margin_left = Inches(0.35)
    tfv.margin_top = Inches(0.3)
    tfv.margin_right = Inches(0.35)

    p = tfv.paragraphs[0]
    p.text = "VERIFIED CORE USER INTERFACE VIEWS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    views_list = [
        ("Cinematic 3D Landing Page", "Three.js WebGL scene with scroll-driven choreography across 5 stages, frosted glass skill capsules, and responsive mobile FOV."),
        ("Personal Dashboard (Home)", "5-metric snapshot bar, active ongoing swaps with direct chat links, pending triage, and recommended match previews."),
        ("Discover Grid & Multi-Filters", "One-user-per-card grid layout with keyword search, category dropdowns, level badges, and mobile slide-over filter drawer."),
        ("Interactive Matches Explorer", "Match score box (0-100%), directional chips ('They Can Teach You' vs 'You Can Teach Them'), and AI toggle."),
        ("Desktop 2-Panel Chat & Meetings", "Full-height workspace with sidebar conversation list, message search, WhatsApp-style date separators, and embedded Jitsi modal."),
        ("Portfolio Media Lightbox", "3-column grid with optimistic upload percentage overlays, 4 emoji reactions with animated bursts, and 5s undo toast.")
    ]
    for vname, vdesc in views_list:
        pv = tfv.add_paragraph()
        pv.text = f"• {vname}"
        pv.font.size = Pt(11)
        pv.font.bold = True
        pv.font.color.rgb = INK_MAIN
        pv.space_before = Pt(6)
        pvd = tfv.add_paragraph()
        pvd.text = f"   {vdesc}"
        pvd.font.size = Pt(9.5)
        pvd.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 13: TESTING AND RESULTS
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide13, "Testing & Verification — 53 Automated Tests", "TEST SUITE & RESULTS")

    suites = [
        ("Recommendation Engine Suite", "15 Tests • 100% Passed",
         "Verifies canonical skill normalization ('JS' <-> 'JavaScript'), false-equivalence guards ('Java' != 'JavaScript'), progressive diminishing returns scoring (72% -> 92% -> 100%), mutual bonuses, and display name preservation."),
        ("Video Meetings & Scheduling Suite", "10 Tests • 100% Passed",
         "Validates accepted swap access permissions, unauthorized blocking (403), instant and scheduled session generation, participant access control, cancellation events, and reminder job eligibility."),
        ("Notification Center Suite", "12 Tests • 100% Passed",
         "Tests real-time Socket.IO room dispatch, user isolation, unread count accuracy, pagination, individual and bulk mark-as-read updates, swap/rating producers, and Map deduplication."),
        ("Rating & Reputation Suite", "4 Tests • 100% Passed",
         "Ensures rating links to exact swap, populates exchanged skill pair metadata, verifies database compound unique index preventing duplicates, and checks server-side MongoDB aggregation recalculating avgRating."),
        ("Gemini Flash AI Matching Suite", "12 Tests • 100% Passed",
         "Verifies default 0-call deterministic mode vs user-triggered ?ai=true execution, exact 70/30 hybrid score math, score bounds (0-100), anti-hallucination candidate validation, and resilient fallback on timeout/429.")
    ]

    for i, (title, status, desc) in enumerate(suites):
        y = Inches(1.55 + i * 1.08)
        c = add_card(slide13, Inches(0.8), y, Inches(11.7), Inches(0.98))
        tfc = c.text_frame
        tfc.word_wrap = True
        tfc.margin_left = Inches(0.3)
        tfc.margin_top = Inches(0.14)
        tfc.margin_right = Inches(0.3)

        p = tfc.paragraphs[0]
        p.text = f"✅ {title} — "
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PINE_DARK

        # Append status in green
        run = p.add_run()
        run.text = status
        run.font.bold = True
        run.font.color.rgb = EMERALD

        p2 = tfc.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = INK_MUTED
        p2.space_before = Pt(2)

    # Bottom Build Verification Bar
    build_box = add_card(slide13, Inches(0.8), Inches(6.15), Inches(11.7), Inches(0.75), bg_color=PINE_LIGHT, border_color=EMERALD)
    tfb = build_box.text_frame
    tfb.word_wrap = True
    tfb.margin_left = Inches(0.3)
    tfb.margin_top = Inches(0.12)
    pb = tfb.paragraphs[0]
    pb.text = "📦 FRONTEND BUILD VERIFICATION: Vite 8 production build (`npm run build`) compiles cleanly with 0 errors and 0 type warnings."
    pb.font.size = Pt(10.5)
    pb.font.bold = True
    pb.font.color.rgb = PINE_DARK

    # =========================================================================
    # SLIDE 14: FUTURE ENHANCEMENTS
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide14, "Current Limitations & Future Enhancements", "ROADMAP & SCALABILITY")

    # Left: Implemented vs Limitations
    c_curr = add_card(slide14, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.3))
    tfc = c_curr.text_frame
    tfc.word_wrap = True
    tfc.margin_left = Inches(0.35)
    tfc.margin_top = Inches(0.3)
    tfc.margin_right = Inches(0.35)

    p = tfc.paragraphs[0]
    p.text = "CURRENT SCOPE & HONEST LIMITATIONS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PINE_DARK

    limits = [
        ("No Email Verification / Password Reset", "Users register directly with email/password; forgot-password recovery is out of mini-project scope."),
        ("Single-Instance Socket.IO", "WebSockets operate in-memory on a single Node process; horizontal scaling across multiple servers will require a Redis adapter."),
        ("In-App Only Notifications", "Notifications are stored in MongoDB and pushed over live WebSockets; native mobile push notifications (FCM/APNs) are not yet integrated."),
        ("No Admin UI Panel", "Portfolio moderation reports are logged to the database with categorical reasons, but manual review requires direct DB access."),
        ("Render Free-Tier Spin-Down", "Backend may experience 30–60 second cold starts after inactivity on free-tier hosting.")
    ]
    for h, d in limits:
        ph = tfc.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = INK_MAIN
        ph.space_before = Pt(6)
        pd = tfc.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = INK_MUTED

    # Right: Future Enhancements
    c_fut = add_card(slide14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), bg_color=CARD_BG_ALT)
    tff = c_fut.text_frame
    tff.word_wrap = True
    tff.margin_left = Inches(0.35)
    tff.margin_top = Inches(0.3)
    tff.margin_right = Inches(0.35)

    p = tff.paragraphs[0]
    p.text = "FUTURE ENHANCEMENT ROADMAP"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    enhancements = [
        ("Redis Multi-Server Scaling", "Add @socket.io/redis-adapter to scale Socket.IO across multiple cluster nodes and share session states."),
        ("Transactional Email & SMS Service", "Integrate Resend or SendGrid for account verification, password resets, and session reminder alerts."),
        ("Multi-Party Skill Swaps", "Extend the bilateral swap engine to support triangular or group barter circles (A teaches B, B teaches C, C teaches A)."),
        ("Dedicated Admin Moderation Suite", "Build a role-based dashboard for community managers to inspect flagged portfolio media and resolve disputes."),
        ("React Native Mobile Application", "Port the React 19 client to React Native for native iOS & Android apps with background push notifications."),
        ("Collaborative AI Filtering", "Train machine learning models on historical swap completions to predict partner chemistry beyond skill keywords.")
    ]
    for h, d in enhancements:
        ph = tff.add_paragraph()
        ph.text = f"✨ {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = INK_MAIN
        ph.space_before = Pt(6)
        pd = tff.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = INK_MUTED

    # =========================================================================
    # SLIDE 15: CONCLUSION & Q&A
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    # Dark luxury closing slide
    bg15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = CHARCOAL_BG
    bg15.line.fill.background()

    top_bar15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar15.fill.solid()
    top_bar15.fill.fore_color.rgb = EMERALD
    top_bar15.line.fill.background()

    # Left side: Key Takeaways
    c_takeaways = add_card(slide15, Inches(1.0), Inches(1.4), Inches(6.8), Inches(5.1), bg_color=RGBColor(24, 28, 26), border_color=PINE_DARK)
    tft = c_takeaways.text_frame
    tft.word_wrap = True
    tft.margin_left = Inches(0.4)
    tft.margin_top = Inches(0.4)
    tft.margin_right = Inches(0.4)

    p = tft.paragraphs[0]
    p.text = "PROJECT SUMMARY & TAKEAWAYS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    takeaways = [
        ("Democratic Learning Model", "Eliminated financial barriers to education through an authentic bilateral barter marketplace."),
        ("Complete End-to-End Engineering", "Shipped all 12 planned phases including 3D WebGL scenes, full dark mode, real-time chat, Jitsi video, and portfolio showcases."),
        ("Responsible AI Architecture", "Implemented a 70/30 hybrid anchor where AI augments deterministic matching rather than hallucinates false compatibility."),
        ("Robust Codebase & Documentation", "Backed by 53 verified automated tests, 12 thorough documentation manuals (164 KB), and zero invented features.")
    ]
    for h, d in takeaways:
        ph = tft.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(12)
        ph.font.bold = True
        ph.font.color.rgb = SURFACE_WHITE
        ph.space_before = Pt(10)
        pd = tft.add_paragraph()
        pd.text = d
        pd.font.size = Pt(11)
        pd.font.color.rgb = RGBColor(160, 180, 170)

    # Right side: Thank You & Q&A
    c_qa = add_card(slide15, Inches(8.2), Inches(1.4), Inches(4.1), Inches(5.1), bg_color=PINE_DARK, border_color=EMERALD)
    tfq = c_qa.text_frame
    tfq.word_wrap = True
    tfq.margin_left = Inches(0.4)
    tfq.margin_top = Inches(0.6)
    tfq.margin_right = Inches(0.4)

    pq1 = tfq.paragraphs[0]
    pq1.text = "Thank You!"
    pq1.font.size = Pt(36)
    pq1.font.bold = True
    pq1.font.color.rgb = SURFACE_WHITE
    pq1.font.name = "Georgia"

    pq2 = tfq.add_paragraph()
    pq2.text = "Questions & Answers"
    pq2.font.size = Pt(20)
    pq2.font.bold = True
    pq2.font.color.rgb = EMERALD
    pq2.space_before = Pt(8)

    pq3 = tfq.add_paragraph()
    pq3.text = "We are open for evaluation, demonstration, and viva questions."
    pq3.font.size = Pt(12)
    pq3.font.color.rgb = RGBColor(210, 230, 220)
    pq3.space_before = Pt(16)

    pq4 = tfq.add_paragraph()
    pq4.text = "Repository & Documentation:\ngithub.com/ShaikMahaboobbasha2005/SkillSwap"
    pq4.font.size = Pt(11)
    pq4.font.color.rgb = RGBColor(180, 210, 195)
    pq4.space_before = Pt(24)

    # Save presentation
    output_dir = os.path.join(base_dir, "presentation")
    os.makedirs(output_dir, exist_ok=True)
    out_path_pres = os.path.join(output_dir, "SkillSwap_Project_Presentation.pptx")
    out_path_root = os.path.join(base_dir, "SkillSwap_Project_Presentation.pptx")

    prs.save(out_path_pres)
    prs.save(out_path_root)
    print(f"[SUCCESS] Presentation successfully created and saved to:\n  1. {out_path_pres}\n  2. {out_path_root}")

if __name__ == "__main__":
    create_presentation()
